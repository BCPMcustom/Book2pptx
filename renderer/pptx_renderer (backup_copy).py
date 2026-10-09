import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN



class PowerPointRenderer:


    def render(self, presentation_spec):
        prs = Presentation("renderer/templates/Template.pptx")
        slide_layout = prs.slide_layouts[1]

        txt_width = Inches(9.6)
        

        slide = prs.slides[0]
        shapes = slide.shapes
        titleBar = shapes.add_shape(MSO_SHAPE.RECTANGLE,       #This defines the MAX size of the title bar
                                Inches(0.7), Inches(0.75), txt_width, Inches(1.5))        
        fill = titleBar.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(155, 155, 155)

        line = titleBar.line
        line.color.rgb = RGBColor(155, 155, 155)
        line.color.brightness = -0.15 # 15% darker
        line.width = Pt(2.5)


        txBoxTitle = slide.shapes.add_textbox(Inches(0.7), Inches(0.75), txt_width, Inches(1.5))
        tf = txBoxTitle.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()

        font = run.font
        font.size = Pt(42)
        font.color.rgb = RGBColor(255, 255, 255)

        run.text = presentation_spec.title



        slide = prs.slides.add_slide(slide_layout)
        shapes = slide.shapes
        closingTitleBar = shapes.add_shape(MSO_SHAPE.RECTANGLE,       #This defines the MAX size of the title bar
                                Inches(0.7), Inches(1.5), txt_width, Inches(5))
        fill = closingTitleBar.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(155, 155, 155)  

        line = closingTitleBar.line
        line.color.rgb = RGBColor(155, 155, 155)
        line.color.brightness = -0.15 # 15% darker
        line.width = Pt(2.5) 

        txBoxClosing = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), txt_width, Inches(5))
        tf = txBoxClosing.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = "\n" + presentation_spec.closing_title

        font = run.font
        font.size = Pt(55)
        font.color.rgb = RGBColor(255, 255, 255)



        print("\n\nDEBUG SECTION\n\n")


        return prs

        
