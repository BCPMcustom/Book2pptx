"""PowerPoint rendering for book2pptx.

!!! VERSION 1.0 code -> refactored by Claude from Ben's code !!!

Turns a ``PresentationSpec`` into a python-pptx ``Presentation``. Two modes:

* Plain (default): the default python-pptx presentation with black text.
* Template: ``renderer/templates/Template.pptx`` with grey title/body blocks
  drawn behind white text.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from parser.models.presentation_spec import PresentationSpec, SlideSpec

TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "Template.pptx"

# Geometry. The template slides are wider than the default 4:3 slide.
LEFT = Inches(0.7)
BLOCK_WIDTH = Inches(8.6)
BLOCK_WIDTH_TEMPLATE = Inches(9.6)
TITLE_TOP = Inches(0.75)
TITLE_HEIGHT = Inches(1.5)
BODY_TOP = Inches(2.5)
BODY_HEIGHT = Inches(4.5)
CLOSING_TOP = Inches(0.9)
CLOSING_HEIGHT = Inches(5.5)

# Images start at IMAGE_ORIGIN and each further image shifts by IMAGE_STAGGER.
IMAGE_ORIGIN = Inches(1)
IMAGE_STAGGER = Inches(0.5)

# Typography.
TITLE_FONT_SIZE = Pt(42)
BODY_FONT_SIZE = Pt(22)
LIGHT_TEXT = RGBColor(255, 255, 255)
DARK_TEXT = RGBColor(0, 0, 0)

# Block styling (template mode).
BLOCK_FILL = RGBColor(155, 155, 155)
BLOCK_OUTLINE_WIDTH = Pt(2.5)
BLOCK_OUTLINE_BRIGHTNESS = -0.15  # Outline is 15% darker than the fill.


class PowerPointRenderer:
    """Render a ``PresentationSpec`` to a PowerPoint deck."""

    def __init__(self, use_template: bool = False) -> None:
        self.use_template = use_template
        self.block_width = BLOCK_WIDTH_TEMPLATE if use_template else BLOCK_WIDTH
        self.text_color = LIGHT_TEXT if use_template else DARK_TEXT
        # Slide layout used for every slide: index 1 in the template, index 0
        # in the default presentation.
        self.layout_index = 1 if use_template else 0

    # -- Shared helpers ------------------------------------------------------

    def _add_block(self, slide, top: int, height: int) -> None:
        """Draw a filled, outlined rectangle behind the text."""
        block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, LEFT, top, self.block_width, height)
        block.fill.solid()
        block.fill.fore_color.rgb = BLOCK_FILL
        block.line.color.rgb = BLOCK_FILL
        block.line.color.brightness = BLOCK_OUTLINE_BRIGHTNESS
        block.line.width = BLOCK_OUTLINE_WIDTH

    def _add_text_frame(self, slide, top: int, height: int):
        """Add a word-wrapped text box spanning the block width."""
        text_box = slide.shapes.add_textbox(LEFT, top, self.block_width, height)
        frame = text_box.text_frame
        frame.word_wrap = True
        return frame

    def _write(self, paragraph, text: str, size, align) -> None:
        """Write ``text`` into ``paragraph`` as a single styled run."""
        paragraph.alignment = align
        run = paragraph.add_run()
        run.font.size = size
        run.font.color.rgb = self.text_color
        run.text = text

    def _write_lines(self, frame, lines: list[str], size, align) -> None:
        """Write each line as its own paragraph, reusing the frame's first one."""
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            self._write(paragraph, line, size, align)

    # -- Blocks and text -----------------------------------------------------

    def render_title_block(self, slide) -> None:
        self._add_block(slide, TITLE_TOP, TITLE_HEIGHT)

    def render_body_block(self, slide) -> None:
        self._add_block(slide, BODY_TOP, BODY_HEIGHT)

    def render_closing_block(self, slide) -> None:
        self._add_block(slide, CLOSING_TOP, CLOSING_HEIGHT)

    def render_title(self, slide, title: str) -> None:
        frame = self._add_text_frame(slide, TITLE_TOP, TITLE_HEIGHT)
        self._write(frame.paragraphs[0], title, TITLE_FONT_SIZE, PP_ALIGN.CENTER)

    def render_closing(self, slide, title: str) -> None:
        frame = self._add_text_frame(slide, CLOSING_TOP, CLOSING_HEIGHT)
        self._write(frame.paragraphs[0], title, TITLE_FONT_SIZE, PP_ALIGN.CENTER)

    def render_bullets(self, slide, slide_spec: SlideSpec) -> None:
        if not slide_spec.bullets:
            return
        frame = self._add_text_frame(slide, BODY_TOP, BODY_HEIGHT)
        lines = [f"- {bullet}" for bullet in slide_spec.bullets]
        self._write_lines(frame, lines, BODY_FONT_SIZE, PP_ALIGN.LEFT)

    def render_callouts(self, slide, slide_spec: SlideSpec) -> None:
        if not slide_spec.callouts:
            return
        frame = self._add_text_frame(slide, BODY_TOP, BODY_HEIGHT)
        self._write_lines(frame, slide_spec.callouts, BODY_FONT_SIZE, PP_ALIGN.LEFT)

    def render_images(self, slide, slide_spec: SlideSpec) -> None:
        for index, image_path in enumerate(slide_spec.images):
            offset = IMAGE_ORIGIN + index * IMAGE_STAGGER
            slide.shapes.add_picture(str(Path(image_path).resolve()), offset, offset)

    # -- Deck assembly -------------------------------------------------------

    def render(self, presentation_spec: PresentationSpec) -> Presentation:
        if self.use_template:
            if not TEMPLATE_PATH.is_file():
                raise FileNotFoundError(f"Template not found at '{TEMPLATE_PATH}'.")
            prs = Presentation(str(TEMPLATE_PATH))
        else:
            prs = Presentation()
        layout = prs.slide_layouts[self.layout_index]

        # Opening slide. The template ships with its title slide already in place.
        if self.use_template:
            slide = prs.slides[0]
            self.render_title_block(slide)
        else:
            slide = prs.slides.add_slide(layout)
        self.render_title(slide, presentation_spec.title)

        for slide_spec in presentation_spec.slides:
            slide = prs.slides.add_slide(layout)
            if self.use_template:
                self.render_title_block(slide)
                self.render_body_block(slide)
            self.render_title(slide, slide_spec.title)
            self.render_callouts(slide, slide_spec)
            self.render_bullets(slide, slide_spec)
            self.render_images(slide, slide_spec)

        # Closing slide.
        slide = prs.slides.add_slide(layout)
        if self.use_template:
            self.render_closing_block(slide)
        self.render_closing(slide, presentation_spec.closing_title)
        return prs
